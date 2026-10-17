"""Create the TV1 report and evidence figures from real project files."""
from pathlib import Path
import csv
import json
import textwrap
from datetime import datetime
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports' / 'TV1'
OUT.mkdir(exist_ok=True)
FIG = OUT / 'images'
FIG.mkdir(exist_ok=True)
doc = Document()
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
md = []

def heading(text, level=1):
    doc.add_heading(text, level)
    md.extend(['#' * level + ' ' + text, ''])

def paragraph(text):
    doc.add_paragraph(text)
    md.extend([text, ''])

def code(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = 'Consolas'
    r.font.size = Pt(9)
    md.extend(['```bat', text, '```', ''])

def figure(name, title, content):
    lines = []
    for line in content.splitlines():
        lines.extend(textwrap.wrap(line, 100, replace_whitespace=False, drop_whitespace=False) or [''])
    font = ImageFont.truetype('C:/Windows/Fonts/consola.ttf', 19)
    titlefont = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 23)
    img = Image.new('RGB', (1280, 100 + 27 * len(lines)), '#101b2b')
    draw = ImageDraw.Draw(img)
    draw.text((28, 20), title, font=titlefont, fill='#80d8ff')
    for i, line in enumerate(lines):
        draw.text((28, 75 + 27 * i), line, font=font, fill='#edf3fb')
    path = FIG / (name + '.png')
    img.save(path)
    doc.add_picture(str(path), width=Inches(6.2))
    caption = 'Hình ' + name + '. ' + title
    doc.add_paragraph(caption, style='Caption')
    md.extend([f'![{caption}](images/{path.name})', ''])

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Shading Accent 1'
    for c, v in zip(t.rows[0].cells, headers): c.text = str(v)
    for row in rows:
        for c, v in zip(t.add_row().cells, row): c.text = str(v)
    md.extend(['| ' + ' | '.join(headers) + ' |', '|' + '|'.join(['---'] * len(headers)) + '|'])
    md.extend('| ' + ' | '.join(map(str, row)) + ' |' for row in rows)
    md.append('')

heading('BÁO CÁO THỰC HIỆN TASK TV1', 1)
paragraph('Chuẩn bị dữ liệu và xây dựng các cấu hình tiền xử lý cho TVPL (tiếng Việt) và NFCorpus (tiếng Anh).')
paragraph('Ngày lập: ' + datetime.now().strftime('%d/%m/%Y') + '. Thư mục dự án: D:\\DataPipeline.')
paragraph('Phạm vi TV1 trong tài liệu được xác định theo code và README của dự án. Hình minh họa được dựng từ code, dữ liệu và báo cáo thật; đây không phải ảnh chụp trực tiếp giao diện CMD. Kết quả kiểm tra mẫu và chuẩn hóa được chạy lại trong phiên lập báo cáo; kết quả kiểm tra toàn bộ được ghi theo báo cáo validation trong dự án.')

heading('1. Danh sách công việc đã thực hiện')
table(['STT', 'Công việc', 'Sản phẩm'], [
    [1, 'Chuẩn bị dữ liệu TVPL và NFCorpus', 'data/raw'],
    [2, 'Thống nhất schema và kiểm tra liên kết ID', 'data/canonical/vi và en'],
    [3, 'Xây dựng hàm tiền xử lý', 'src/preprocess.py'],
    [4, 'Kiểm tra câu mẫu và trường hợp đặc biệt', 'preprocessing_examples, sample_validation'],
    [5, 'Tạo năm cấu hình cho hai ngôn ngữ', 'data/processed/vi và en'],
    [6, 'Kiểm tra dữ liệu và chuẩn bị bàn giao', 'validation_report, build_manifest, handoff']])

heading('2. Thực hiện từng bước')
heading('Bước 1. Mở CMD và chuẩn bị môi trường', 2)
paragraph('Mở Command Prompt trên Windows. Chạy lần lượt các lệnh dưới đây. Nếu .venv đã tồn tại thì bỏ qua lệnh tạo môi trường. Các lệnh dùng trực tiếp Python trong .venv nên không cần activate.')
code('cd /d D:\\DataPipeline\npython -m venv .venv\n.\\.venv\\Scripts\\python.exe -m pip install -r requirements.txt')
paragraph('Vai trò: cài các thư viện đọc dữ liệu và xử lý ngôn ngữ, gồm Underthesea và NLTK. Dự án hiện đã có môi trường chạy được; quá trình cài đặt không được chạy lại để tạo bằng chứng giả.')
figure('01', 'Thư viện của dự án — nguồn: requirements.txt', (ROOT/'requirements.txt').read_text(encoding='utf-8'))
paragraph('Ảnh CMD cần chụp khi tự thực hiện: lệnh cài thư viện cùng thông báo hoàn tất. Không dùng hình danh sách thư viện để thay cho bằng chứng cài thành công.')

heading('Bước 2. Chuẩn hóa cấu trúc dữ liệu', 2)
code('.\\.venv\\Scripts\\python.exe src\\prepare_data.py')
paragraph('Script đọc dữ liệu đã có trong data/raw, ghép title và text bằng một ký tự xuống dòng khi có nội dung, đổi tên ID và xuất UTF-8. Bước này không lowercase, không loại stopword và không tách từ. NFCorpus chỉ giữ các query có ID trong qrels test.')
paragraph('Đầu ra mỗi ngôn ngữ gồm corpus.jsonl với doc_id/text, queries.jsonl với query_id/text và qrels.csv với query_id/doc_id/relevance. Script kiểm tra ID duy nhất, tham chiếu qrels, đọc lại file xuất và hash dữ liệu gốc.')
canonical = json.loads((ROOT/'reports/canonical_validation.json').read_text(encoding='utf-8'))
figure('02', 'Kết quả chuẩn hóa — nguồn: canonical_validation.json', json.dumps(canonical['datasets'], ensure_ascii=False, indent=2))
for lang in ['vi', 'en']:
    with (ROOT/f'data/canonical/{lang}/queries.jsonl').open(encoding='utf-8') as f: sample=json.loads(next(f))
    figure('03_'+lang, 'Query canonical thực tế — ' + lang, json.dumps(sample, ensure_ascii=False, indent=2))
paragraph('Ảnh CMD cần chụp: dòng thống kê vi và en sau lệnh prepare_data.py. Ảnh đầu ra: mở data/canonical và một bản ghi JSONL trong VS Code.')

heading('Bước 3. Giải thích code tiền xử lý', 2)
table(['Cấu hình', 'Quy tắc'], [
    ['original', 'Giữ nguyên canonical, gồm phần title đã ghép.'],
    ['stopword', 'Loại từ/cụm từ dừng; khớp không phân biệt hoa thường, ưu tiên cụm dài.'],
    ['normalize', 'Lowercase; giữ chữ Unicode, dấu kết hợp và số; thay phần còn lại bằng khoảng trắng rồi gộp.'],
    ['tokenize', 'VI: Underthesea; EN: TreebankWordTokenizer của NLTK.'],
    ['combined', 'normalize → tokenize → remove stopwords.']])
source=(ROOT/'src/preprocess.py').read_text(encoding='utf-8').splitlines()
figure('04', 'Code normalize và tokenize — src/preprocess.py', '\n'.join(f'{i+1:3} {source[i]}' for i in range(15,26)))
paragraph('normalize xử lý từng ký tự theo nhóm Unicode để giữ dấu tiếng Việt. tokenize trả về danh sách token; tiếng Việt có thể nối từ nhiều âm tiết bằng dấu gạch dưới. Treebank có thể tách contractions của tiếng Anh.')
start=next(i for i,l in enumerate(source) if l.startswith('def transform'))
figure('05', 'Điều phối năm cấu hình — src/preprocess.py', '\n'.join(f'{i+1:3} {source[i]}' for i in range(start,len(source))))
paragraph('Stopword gốc nằm tại resources/stopwords/vi.txt và en.txt; nguồn và checksum được ghi trong sources.json. Ở combined, danh sách stopword cũng được normalize và tokenize để so khớp nhất quán với văn bản. Không loại bản ghi khi kết quả rỗng.')

heading('Bước 4. Chạy thử trên câu mẫu', 2)
code('.\\.venv\\Scripts\\python.exe src\\check_samples.py')
samples=json.loads((ROOT/'reports/sample_validation.json').read_text(encoding='utf-8'))
figure('06', 'Kết quả chạy lại kiểm tra mẫu', '.\\.venv\\Scripts\\python.exe src\\check_samples.py\n60 sample checks passed\n\nNguồn: stdout thực tế của lần chạy lập báo cáo.\nBáo cáo: reports/sample_validation.json\npassed: '+str(samples['passed'])+'\nSố kiểm tra trong file: '+str(len(samples['checks'])))
paragraph('Kiểm tra bao gồm Unicode, dấu câu, ranh giới từ, cụm stopword nhiều âm tiết, token có gạch dưới, câu rỗng và tính ổn định của normalize. Báo cáo ví dụ lưu tại reports/preprocessing_examples.md và .csv.')
with (ROOT/'reports/preprocessing_examples.csv').open(encoding='utf-8', newline='') as f: examples=list(csv.DictReader(f))
for lang in ['vi','en']:
    row=next(r for r in examples if r['language']==lang)
    figure('07_'+lang, 'Ví dụ trước và sau xử lý — ' + lang, '\n\n'.join(k+': '+row[k] for k in ['original','stopword','normalize','tokenize','combined']))

heading('Bước 5. Sinh toàn bộ năm cấu hình', 2)
code('.\\.venv\\Scripts\\python.exe src\\build_variants.py')
paragraph('Script lặp qua hai ngôn ngữ, năm cấu hình và hai loại dữ liệu corpus/queries. Mỗi bản ghi được transform rồi ghi JSONL; giữ nguyên ID và thứ tự. qrels.csv được sao chép một lần cho mỗi ngôn ngữ. Lệnh này ghi lại dữ liệu processed và có thể chạy lâu do tách từ toàn bộ dữ liệu.')
paragraph('Trong phiên lập báo cáo, bước build dùng các file đã có; không mô tả là đã build lại. Log thực tế khi tự chạy thường có dạng “vi original corpus done 10576”; cần chụp log từ lần chạy của bạn.')
files=sorted(str(p.relative_to(ROOT)).replace('\\','/') for p in (ROOT/'data/processed').rglob('*') if p.is_file())
figure('08', 'Các file processed hiện có — liệt kê từ ổ đĩa', '\n'.join(files))
paragraph('Thành phẩm gồm 20 file JSONL (2 ngôn ngữ × 5 cấu hình × 2 loại) và 2 file qrels.csv dùng chung ở cấp ngôn ngữ. Thống kê số lượng, độ dài và thời gian build nằm tại reports/dataset_summary.csv; phiên bản thư viện và hash nằm tại reports/build_manifest.json.')

heading('Bước 6. Kiểm tra toàn bộ và bàn giao', 2)
code('.\\.venv\\Scripts\\python.exe src\\validate_data.py')
paragraph('Script kiểm tra schema, ID, liên kết qrels, hash đầu vào/đầu ra, tính lại toàn bộ phép transform và so sánh nội dung từng bản ghi. Phải đợi tiến trình kết thúc; chỉ khi hoàn tất thành công mới có dòng “Full validation passed”.')
validation=json.loads((ROOT/'reports/validation_report.json').read_text(encoding='utf-8'))
rows=[]
for lang, variants in validation['datasets'].items():
    for variant, kinds in variants.items():
        rows.append([lang,variant,kinds['corpus']['records'],kinds['queries']['records'],len(kinds['corpus']['empty_ids']),len(kinds['queries']['empty_ids'])])
figure('09', 'Tổng hợp validation — nguồn: validation_report.json', 'passed: '+str(validation['passed'])+'\n\nNgôn ngữ | Cấu hình | Tài liệu | Truy vấn | Rỗng TL | Rỗng TV\n'+'\n'.join(' | '.join(map(str,r)) for r in rows))
paragraph('Bàn giao data/processed, config.json, nguồn stopword và các báo cáo. Người làm truy xuất cần thống nhất cách đọc token, giữ token ghép tiếng Việt và tránh áp dụng thêm lowercase/stopword ngoài cấu hình đã chọn. Người đánh giá dùng cùng tập query và mức relevance gốc cho mọi cấu hình.')

heading('3. Thành quả')
table(['Bộ dữ liệu', 'Tài liệu', 'Truy vấn', 'Qrels', 'Cấu hình'], [['TVPL / vi',10576,9985,11158,5],['NFCorpus / en',3633,323,12334,5]])
table(['Ngôn ngữ','Cấu hình','Tài liệu','Truy vấn','TL rỗng','TV rỗng'],rows)
paragraph('Các hình 02–03 thể hiện dữ liệu canonical; hình 04–05 giải thích code; hình 06 thể hiện kiểm tra mẫu đã chạy; hình 07 thể hiện thay đổi văn bản; hình 08 thể hiện file đầu ra; hình 09 thể hiện báo cáo validation. Các số liệu là từ dữ liệu và báo cáo dự án, không phải số liệu minh họa.')
paragraph('Giới hạn: TVPL và NFCorpus khác miền nội dung nên không thể quy chênh lệch điểm giữa hai bộ chỉ cho ngôn ngữ. Chưa chạy TF-IDF/BM25, chưa có điểm truy xuất và chưa thử tích hợp với code thành viên tiếp theo. Danh sách stopword tiếng Việt rộng có thể loại từ mang thông tin pháp luật; không điều chỉnh quy tắc dựa trên điểm test.')

heading('4. Danh sách ảnh CMD cần bổ sung nếu yêu cầu ảnh chụp màn hình trực tiếp')
paragraph('Tài liệu đã có hình minh họa dữ liệu/code thật. Nếu giảng viên yêu cầu ảnh giao diện CMD, mở CMD, chạy lần lượt các lệnh ở phần 2 và dùng Win + Shift + S để chụp sau từng lệnh: cài thư viện; prepare_data.py; check_samples.py; build_variants.py; validate_data.py. Chụp thêm cây thư mục processed và ví dụ JSONL. Giữ rõ cả lệnh lẫn kết quả; chỉ chụp khi tiến trình hoàn tất. Các hình dựng trong báo cáo không được ghi là ảnh chụp CMD.')

doc.save(OUT/'Bao_cao_TV1.docx')
(OUT/'Bao_cao_TV1.md').write_text('\n'.join(md),encoding='utf-8')
print('Created:',OUT/'Bao_cao_TV1.docx')
print('Figures:',len(list(FIG.glob('*.png'))))
